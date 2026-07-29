"""Text extraction, one function per file format.

A parser turns a file into plain text plus whatever metadata the format itself
knows — a PDF's page count, an HTML document's `<title>`. Everything downstream
(chunking, embedding, citation) operates on text alone, so supporting a new format
is a new function and one entry in `PARSERS`; no other module changes.

Two deliberate properties:

* **Lazy imports.** Formats needing a third-party library import it inside the
  parser and raise `MissingDependencyError`. A corpus containing one PDF still
  ingests completely when `pypdf` is absent — the loader records that single file
  as a failure and continues.
* **No rewriting.** Parsers extract; they do not summarise, translate or reflow.
  Chunk text is quoted back to users as citation evidence, so text the pipeline
  invented would make that evidence a forgery.
"""

from __future__ import annotations

from collections.abc import Callable, Mapping
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from ..errors import MissingDependencyError, ParseError


@dataclass(frozen=True, slots=True)
class ParsedContent:
    """The result of extracting one file.

    `title` is set only when the format declares one (an HTML `<title>`, a DOCX
    core property). It takes precedence over the title the loader would otherwise
    derive, because a format's own title is authoritative and citation titles are
    what a user reads when deciding whether to trust an answer.
    """

    text: str
    title: str | None = None
    metadata: Mapping[str, Any] = field(default_factory=dict)


type Parser = Callable[[Path, str], ParsedContent]


def parse_text(path: Path, encoding: str) -> ParsedContent:
    """Read a plain-text format (Markdown, reStructuredText, plain text)."""
    try:
        return ParsedContent(text=path.read_text(encoding=encoding))
    except UnicodeDecodeError as exc:
        raise ParseError(f"{path.name} is not valid {encoding} text: {exc}") from exc


def parse_html(path: Path, encoding: str) -> ParsedContent:
    """Extract readable text from HTML using the standard library.

    No dependency: `html.parser` is sufficient to drop markup, and a heavier
    library would earn its place only if we needed layout-aware extraction such as
    main-content detection.
    """
    extractor = _HtmlTextExtractor()
    extractor.feed(path.read_text(encoding=encoding, errors="replace"))
    extractor.close()
    return ParsedContent(
        text=extractor.text(),
        title=extractor.title,
        metadata={"content_type": "text/html"},
    )


def parse_pdf(path: Path, _encoding: str) -> ParsedContent:
    """Extract text from a PDF, page by page.

    Page boundaries become blank lines so the chunker can split on them like any
    other paragraph break, and `page_count` is retained because "which page?" is
    the first question asked of a PDF citation.
    """
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise MissingDependencyError("pdf", "pypdf", "documents") from exc

    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            # An empty-password decrypt covers PDFs that are "protected" only
            # against editing, which is most of them in practice.
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ParseError(f"{path.name} is password-protected.") from exc
        pages = [page.extract_text() or "" for page in reader.pages]
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"Could not read PDF {path.name}: {exc}") from exc

    # PDF extractors pad lines out to the page width. That padding is a layout
    # artefact with no meaning, and left in place it pollutes both the embedding
    # and the text quoted back in a citation. Only trailing space is removed; no
    # word is altered.
    stripped = ["\n".join(line.rstrip() for line in page.splitlines()) for page in pages]
    text = "\n\n".join(page.strip() for page in stripped if page.strip())
    if not text.strip():
        # A scanned PDF parses without error and yields nothing. Saying so is far
        # more useful than indexing an empty document.
        raise ParseError(
            f"{path.name} contains no extractable text. It is most likely a scan; "
            f"OCR is not part of the ingestion pipeline."
        )
    return ParsedContent(
        text=text,
        metadata={"content_type": "application/pdf", "page_count": len(reader.pages)},
    )


def parse_docx(path: Path, _encoding: str) -> ParsedContent:
    """Extract paragraphs and table cells from a Word document.

    Tables are included because internal policy documents keep exactly the facts
    people ask about — limits, rates, entitlements — inside them. Cells are joined
    with a tab so a row survives as one line of text.
    """
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise MissingDependencyError("docx", "python-docx", "documents") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ParseError(f"Could not read DOCX {path.name}: {exc}") from exc

    blocks = [para.text.strip() for para in document.paragraphs if para.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append("\t".join(cells))

    title = (document.core_properties.title or "").strip() or None
    return ParsedContent(
        text="\n\n".join(blocks),
        title=title,
        metadata={
            "content_type": (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            "table_count": len(document.tables),
        },
    )


# The supported corpus formats. Adding one is a new function plus a line here.
PARSERS: dict[str, Parser] = {
    ".md": parse_text,
    ".markdown": parse_text,
    ".txt": parse_text,
    ".rst": parse_text,
    ".html": parse_html,
    ".htm": parse_html,
    ".pdf": parse_pdf,
    ".docx": parse_docx,
}

SUPPORTED_EXTENSIONS: tuple[str, ...] = tuple(sorted(PARSERS))


def parse(path: Path, encoding: str = "utf-8") -> ParsedContent:
    """Extract `path` using the parser registered for its extension.

    Raises:
        ParseError: The extension has no parser, or extraction failed.
        MissingDependencyError: The format's optional library is not installed.
    """
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ParseError(
            f"No parser for {path.suffix!r}. Supported: {', '.join(SUPPORTED_EXTENSIONS)}."
        )
    return parser(path, encoding)


class _HtmlTextExtractor(HTMLParser):
    """Collects visible text, discarding markup and non-content elements."""

    # Their content is code or metadata, never prose a user would want retrieved.
    _SKIP = frozenset({"script", "style", "noscript", "template", "svg"})
    # Elements whose boundaries are paragraph breaks in the extracted text, which
    # is what lets the chunker split an HTML document on its real structure.
    _BLOCK = frozenset(
        {
            "p", "div", "br", "li", "tr", "section", "article", "header", "footer",
            "h1", "h2", "h3", "h4", "h5", "h6", "table", "blockquote", "pre",
        }
    )

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title: str | None = None
        self._parts: list[str] = []
        self._skip_depth = 0
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: Any) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag == "title":
            self._in_title = True
        elif tag in self._BLOCK:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag == "title":
            self._in_title = False
        elif tag in self._BLOCK:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title = (self.title or "") + data.strip()
        elif self._skip_depth == 0:
            self._parts.append(data)

    def text(self) -> str:
        """The collected text, with each block element on its own line."""
        lines = ("".join(self._parts)).splitlines()
        return "\n".join(stripped for line in lines if (stripped := line.strip()))
